"""Master CV import (DOCX/PDF → text for review and AI evidence). Tailored document generation is Phase 4."""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from logidx.models import Document
from logidx.settings import get_settings

MAX_BYTES = 5 * 1024 * 1024
ALLOWED = {".docx": "docx", ".pdf": "pdf"}


def extract_text(data: bytes, kind: str) -> str:
    if kind == "docx":
        from docx import Document as Docx

        doc = Docx(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(dict.fromkeys(cells)))
        return "\n".join(parts)
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def import_master_cv(session: Session, filename: str, data: bytes) -> Document:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("upload a .docx or .pdf file")
    if len(data) > MAX_BYTES:
        raise ValueError("file is larger than 5 MB")
    kind = ALLOWED[suffix]
    if kind == "docx" and not data.startswith(b"PK"):
        raise ValueError("file is not a valid .docx")
    if kind == "pdf" and not data.startswith(b"%PDF"):
        raise ValueError("file is not a valid PDF")
    text = extract_text(data, kind)
    version = (session.scalar(select(func.max(Document.version)).where(Document.kind == "master_cv")) or 0) + 1
    folder = get_settings().data_dir / "documents"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"master_cv_v{version}{suffix}"
    path.write_bytes(data)
    doc = Document(
        kind="master_cv",
        version=version,
        filename=Path(filename).name[:300],
        path=str(path),
        sha256=hashlib.sha256(data).hexdigest(),
        extracted_text=text,
    )
    session.add(doc)
    return doc
